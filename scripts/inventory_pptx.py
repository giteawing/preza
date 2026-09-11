#!/usr/bin/env python3
"""Inventory text and raster objects in the three transferred PPTX parts."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def inspect(path: Path, global_offset: int) -> dict:
    prs = Presentation(path)
    slides = []
    for local_no, slide in enumerate(prs.slides, 1):
        text_shapes = []
        pictures = []
        for index, shape in enumerate(slide.shapes):
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                text_shapes.append({
                    "shape_index": index,
                    "name": shape.name,
                    "text": shape.text,
                    "left": shape.left,
                    "top": shape.top,
                    "width": shape.width,
                    "height": shape.height,
                })
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures.append({
                    "shape_index": index,
                    "name": shape.name,
                    "left": shape.left,
                    "top": shape.top,
                    "width": shape.width,
                    "height": shape.height,
                    "pixels": list(shape.image.size),
                    "extension": shape.image.ext,
                })
        slides.append({
            "global_slide": global_offset + local_no,
            "local_slide": local_no,
            "editable_text_shapes": text_shapes,
            "pictures": pictures,
        })
    return {
        "file": path.name,
        "slide_width": prs.slide_width,
        "slide_height": prs.slide_height,
        "slide_count": len(prs.slides),
        "slides": slides,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="ocr_inventory.json")
    args = parser.parse_args()
    offset = 0
    parts = []
    for filename in ("p1.pptx", "p2.pptx", "p3.pptx"):
        part = inspect(Path(filename), offset)
        parts.append(part)
        offset += part["slide_count"]
    Path(args.output).write_text(
        json.dumps({"total_slides": offset, "parts": parts}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
