#!/usr/bin/env python3
"""Write the complete verified slide transcript into PowerPoint speaker notes."""
from __future__ import annotations
import json, re, sys
from pathlib import Path
from pptx import Presentation

DATA=json.loads(Path('ocr_manual_text.json').read_text(encoding='utf-8'))
CONFIG={
    'p1':('p1_editable.pptx',DATA['p1']),
    'p2':('p2_editable.pptx',DATA['p2']),
    'p3':('p3_editable.pptx',DATA['p3']),
}
CORRECTIONS={
    'царапает ведь ха сердце':'царапает ведь за сердце',
    'Дашрáтх':'Дашратх',
    'Мáнджхи':'Манджхи',
}

def clean(text:str)->str:
    text=text.replace('\x0b','\n').replace('\xa0',' ')
    for bad,good in CORRECTIONS.items(): text=text.replace(bad,good)
    return '\n'.join(re.sub(r'[ \t]+',' ',x).strip() for x in text.splitlines() if x.strip())

def norm(text:str)->str:
    return re.sub(r'[^0-9a-zа-яё]+',' ',text.lower()).strip()

def native_blocks(slide):
    found=[]
    for shape in slide.shapes:
        if not getattr(shape,'has_text_frame',False) or shape.name.startswith('EDITABLE OCR'): continue
        text=clean(shape.text)
        if not text or re.fullmatch(r'\d{1,2}',text): continue
        found.append((shape.top,shape.left,text))
    return [t for _,_,t in sorted(found)]

def transcript(slide, manual):
    parts=[]
    if manual: parts.extend(clean(manual).splitlines())
    joined=norm('\n'.join(parts))
    for block in native_blocks(slide):
        for line in block.splitlines():
            n=norm(line)
            if n and n not in joined:
                parts.append(line); joined=norm('\n'.join(parts))
    # Stable de-duplication after punctuation/whitespace normalization.
    result=[]; seen=set()
    for line in parts:
        n=norm(line)
        if n and n not in seen: result.append(line); seen.add(n)
    return result

def apply(part):
    filename,manual=CONFIG[part]
    prs=Presentation(filename)
    for no,slide in enumerate(prs.slides,1):
        lines=transcript(slide,manual.get(str(no),''))
        frame=slide.notes_slide.notes_text_frame
        frame.clear()
        for i,line in enumerate(lines):
            p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
            p.text=line
    prs.save(filename)

def update_guide():
    rows=['### Полный журнал текста по слайдам','', '| Глобальный слайд | Часть | Текст в заметках |','|---:|---|---|']
    global_no=0
    for part,(filename,_) in CONFIG.items():
        prs=Presentation(filename)
        for local_no,slide in enumerate(prs.slides,1):
            global_no+=1
            text=clean(slide.notes_slide.notes_text_frame.text).replace('|','\\|').replace('\n','<br>')
            rows.append(f'| {global_no} | `{part}`, слайд {local_no} | {text} |')
    guide=Path('OCR_GUIDE.md').read_text(encoding='utf-8')
    start='<!-- OCR-NOTES-START -->'; end='<!-- OCR-NOTES-END -->'
    generated='\n'.join(rows)
    guide=guide[:guide.index(start)+len(start)]+'\n'+generated+'\n'+guide[guide.index(end):]
    Path('OCR_GUIDE.md').write_text(guide,encoding='utf-8')

if __name__=='__main__':
    for part in (sys.argv[1:] or CONFIG):
        if part not in CONFIG: raise SystemExit(f'unknown part: {part}')
        apply(part)
    update_guide()
